import os
import io
import json
import zipfile
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, PatternFill, Border, Side
from reportlab.lib.pagesizes import letter, A4
from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Table, TableStyle, Image, HRFlowable, KeepTogether
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib import colors
from reportlab.lib.units import inch, cm

def generate_excel_bytes(client, wedding_data, theme):
    """Generates a well-formatted Excel workbook for a client wedding dataset."""
    wb = Workbook()
    
    # Sheet 1: Ringkasan & Biodata
    ws1 = wb.active
    ws1.title = "Biodata & Acara"
    ws1.views.sheetView[0].showGridLines = True

    # Styling definitions
    title_font = Font(name="Calibri", size=16, bold=True, color="33443C")
    sec_font = Font(name="Calibri", size=12, bold=True, color="FFFFFF")
    sec_fill = PatternFill(start_color="33443C", end_color="33443C", fill_type="solid")
    sec2_fill = PatternFill(start_color="292929", end_color="292929", fill_type="solid")
    label_font = Font(name="Calibri", size=11, bold=True, color="33443C")
    val_font = Font(name="Calibri", size=11, color="292929")
    thin_border = Border(
        left=Side(style='thin', color='E5D6BA'),
        right=Side(style='thin', color='E5D6BA'),
        top=Side(style='thin', color='E5D6BA'),
        bottom=Side(style='thin', color='E5D6BA')
    )
    header_fill = PatternFill(start_color="F6F2E6", end_color="F6F2E6", fill_type="solid")

    row = 1
    ws1.merge_cells(f"A{row}:D{row}")
    client_title = client['couple_name'] if client['couple_name'] else client['username']
    ws1[f"A{row}"] = f"DATA UNDANGAN PERNIKAHAN: {client_title.upper()}"
    ws1[f"A{row}"].font = title_font
    row += 1
    ws1[f"A{row}"] = f"Status Pesanan: {client['order_status']} | Kelengkapan: {wedding_data['completion_percent']}%"
    ws1[f"A{row}"].font = Font(name="Calibri", size=10, italic=True, color="64748B")
    row += 2

    # Section 1: Tema
    ws1.merge_cells(f"A{row}:D{row}")
    ws1[f"A{row}"] = "1. TEMA & DESAIN UNDANGAN"
    ws1[f"A{row}"].font = sec_font
    ws1[f"A{row}"].fill = sec_fill
    row += 1

    theme_name = wedding_data['theme_name'] if ('theme_name' in wedding_data.keys() and wedding_data['theme_name']) else (theme['name'] if theme else '-')
    fields_theme = [
        ("Nama Tema", theme_name)
    ]
    for lbl, val in fields_theme:
        ws1[f"A{row}"] = lbl
        ws1[f"A{row}"].font = label_font
        ws1[f"A{row}"].fill = header_fill
        ws1.merge_cells(f"B{row}:D{row}")
        ws1[f"B{row}"] = val
        ws1[f"B{row}"].font = val_font
        for c in ['A', 'B', 'C', 'D']:
            ws1[f"{c}{row}"].border = thin_border
        row += 1

    row += 1
    # Section 2: Biodata Pengantin
    ws1.merge_cells(f"A{row}:D{row}")
    ws1[f"A{row}"] = "2. BIODATA MEMPELAI"
    ws1[f"A{row}"].font = sec_font
    ws1[f"A{row}"].fill = sec_fill
    row += 1

    ws1.merge_cells(f"A{row}:B{row}")
    ws1[f"A{row}"] = "Mempelai Pria (Groom)"
    ws1[f"A{row}"].font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    ws1[f"A{row}"].fill = sec2_fill

    ws1.merge_cells(f"C{row}:D{row}")
    ws1[f"C{row}"] = "Mempelai Wanita (Bride)"
    ws1[f"C{row}"].font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    ws1[f"C{row}"].fill = sec2_fill
    row += 1

    bio_pairs = [
        ("Nama Lengkap", wedding_data['groom_name'], "Nama Lengkap", wedding_data['bride_name']),
        ("Nama Panggilan", wedding_data['groom_nickname'], "Nama Panggilan", wedding_data['bride_nickname']),
        ("Anak ke", wedding_data['groom_child_order'], "Anak ke", wedding_data['bride_child_order']),
        ("Nama Ayah", wedding_data['groom_father'], "Nama Ayah", wedding_data['bride_father']),
        ("Nama Ibu", wedding_data['groom_mother'], "Nama Ibu", wedding_data['bride_mother']),
        ("Instagram", wedding_data['groom_instagram'], "Instagram", wedding_data['bride_instagram'])
    ]
    for l1, v1, l2, v2 in bio_pairs:
        ws1[f"A{row}"] = f"{l1}: {v1 or '-'}"
        ws1[f"A{row}"].font = val_font
        ws1.merge_cells(f"A{row}:B{row}")
        
        ws1[f"C{row}"] = f"{l2}: {v2 or '-'}"
        ws1[f"C{row}"].font = val_font
        ws1.merge_cells(f"C{row}:D{row}")
        for c in ['A', 'B', 'C', 'D']:
            ws1[f"{c}{row}"].border = thin_border
        row += 1

    row += 1
    # Section 3: Detail Acara
    ws1.merge_cells(f"A{row}:D{row}")
    ws1[f"A{row}"] = "3. RANGKAIAN ACARA PERNIKAHAN"
    ws1[f"A{row}"].font = sec_font
    ws1[f"A{row}"].fill = sec_fill
    row += 1

    events = json.loads(wedding_data['events_json'] or '[]')
    if not events:
        ws1.merge_cells(f"A{row}:D{row}")
        ws1[f"A{row}"] = "(Belum ada data acara)"
        ws1[f"A{row}"].font = val_font
        row += 1
    else:
        for idx, ev in enumerate(events, 1):
            ws1.merge_cells(f"A{row}:D{row}")
            ws1[f"A{row}"] = f"Acara #{idx}: {ev.get('name', 'Acara')} - {ev.get('date', '')} ({ev.get('start_time', '')} s/d {ev.get('end_time', '')})"
            ws1[f"A{row}"].font = Font(name="Calibri", size=11, bold=True, color="1E293B")
            ws1[f"A{row}"].fill = header_fill
            row += 1

            ws1[f"A{row}"] = "Lokasi"
            ws1[f"A{row}"].font = label_font
            ws1.merge_cells(f"B{row}:D{row}")
            ws1[f"B{row}"] = ev.get('location_name', '')
            row += 1

            ws1[f"A{row}"] = "Alamat"
            ws1[f"A{row}"].font = label_font
            ws1.merge_cells(f"B{row}:D{row}")
            ws1[f"B{row}"] = ev.get('address', '')
            row += 1

            ws1[f"A{row}"] = "Google Maps"
            ws1[f"A{row}"].font = label_font
            ws1.merge_cells(f"B{row}:D{row}")
            ws1[f"B{row}"] = ev.get('maps_url', '')
            row += 1

    # Sheet 2: Cerita, Kado, Undangan & Musik
    ws2 = wb.create_sheet(title="Cerita, Kado & Musik")
    ws2.views.sheetView[0].showGridLines = True
    row = 1

    # Section 4: Love Story
    ws2.merge_cells(f"A{row}:D{row}")
    ws2[f"A{row}"] = "4. KISAH CINTA (LOVE STORY)"
    ws2[f"A{row}"].font = sec_font
    ws2[f"A{row}"].fill = sec_fill
    row += 1

    stories = json.loads(wedding_data['love_stories_json'] or '[]')
    if not stories:
        ws2.merge_cells(f"A{row}:D{row}")
        ws2[f"A{row}"] = "(Belum ada kisah cinta)"
        row += 1
    else:
        for s in stories:
            ws2[f"A{row}"] = f"[{s.get('date_or_year', '-')}] {s.get('title', 'Cerita')}"
            ws2[f"A{row}"].font = label_font
            ws2.merge_cells(f"A{row}:D{row}")
            row += 1
            ws2[f"A{row}"] = s.get('story_text', '')
            ws2[f"A{row}"].font = val_font
            ws2.merge_cells(f"A{row}:D{row}")
            row += 1

    row += 1
    # Section 5: Kado Digital
    ws2.merge_cells(f"A{row}:D{row}")
    ws2[f"A{row}"] = "5. KADO DIGITAL & ALAMAT PENGIRIMAN"
    ws2[f"A{row}"].font = sec_font
    ws2[f"A{row}"].fill = sec_fill
    row += 1

    gifts = json.loads(wedding_data['gifts_json'] or '{}')
    bank_accs = gifts.get('bank_accounts', [])
    for b in bank_accs:
        ws2[f"A{row}"] = f"Bank: {b.get('bank_name')}"
        ws2[f"B{row}"] = f"No: {b.get('account_number')}"
        ws2.merge_cells(f"C{row}:D{row}")
        ws2[f"C{row}"] = f"A.N: {b.get('account_holder')}"
        row += 1

    ewallets = gifts.get('ewallets', [])
    for e in ewallets:
        ws2[f"A{row}"] = f"E-Wallet: {e.get('provider')}"
        ws2[f"B{row}"] = f"No: {e.get('number')}"
        ws2.merge_cells(f"C{row}:D{row}")
        ws2[f"C{row}"] = f"A.N: {e.get('account_holder')}"
        row += 1

    phys = gifts.get('physical_gift', {})
    if phys.get('address'):
        ws2[f"A{row}"] = "Kado Fisik"
        ws2.merge_cells(f"B{row}:D{row}")
        ws2[f"B{row}"] = f"Penerima: {phys.get('recipient')} ({phys.get('phone')}) - {phys.get('address')}"
        row += 1

    row += 1
    # Section 6: Turut Mengundang
    ws2.merge_cells(f"A{row}:D{row}")
    ws2[f"A{row}"] = "6. TURUT MENGUNDANG"
    ws2[f"A{row}"].font = sec_font
    ws2[f"A{row}"].fill = sec_fill
    row += 1

    inviters = json.loads(wedding_data['guest_inviters_json'] or '[]')
    if not inviters:
        ws2.merge_cells(f"A{row}:D{row}")
        ws2[f"A{row}"] = "(Belum diisi)"
        row += 1
    else:
        for idx, inv in enumerate(inviters, 1):
            ws2[f"A{row}"] = f"{idx}."
            ws2.merge_cells(f"B{row}:D{row}")
            ws2[f"B{row}"] = inv.get('name', '')
            row += 1

    row += 1
    # Section 7: Backsound
    ws2.merge_cells(f"A{row}:D{row}")
    ws2[f"A{row}"] = "7. BACKSOUND MUSIK"
    ws2[f"A{row}"].font = sec_font
    ws2[f"A{row}"].fill = sec_fill
    row += 1

    ws2[f"A{row}"] = "Tipe Musik"
    ws2[f"B{row}"] = "Pilihan Preset" if wedding_data['music_choice_type'] == 'preset' else "Unggah Musik Sendiri"
    ws2.merge_cells(f"B{row}:D{row}")
    row += 1

    ws2[f"A{row}"] = "Judul Lagu"
    ws2[f"B{row}"] = wedding_data['music_title'] or '-'
    ws2.merge_cells(f"B{row}:D{row}")
    row += 1

    ws2[f"A{row}"] = "Artis / Keterangan"
    ws2[f"B{row}"] = wedding_data['music_artist'] or '-'
    ws2.merge_cells(f"B{row}:D{row}")

    # Set column widths
    for ws in [ws1, ws2]:
        ws.column_dimensions['A'].width = 24
        ws.column_dimensions['B'].width = 30
        ws.column_dimensions['C'].width = 24
        ws.column_dimensions['D'].width = 30

    buf = io.BytesIO()
    wb.save(buf)
    buf.seek(0)
    return buf.getvalue()


def generate_pdf_bytes(client, wedding_data, theme, base_dir):
    """Generates a professional, print-ready PDF summary of the client wedding biodata."""
    buf = io.BytesIO()
    doc = SimpleDocTemplate(
        buf,
        pagesize=A4,
        rightMargin=36,
        leftMargin=36,
        topMargin=36,
        bottomMargin=36
    )

    styles = getSampleStyleSheet()
    
    # Custom styles
    title_style = ParagraphStyle(
        'DocTitle',
        parent=styles['Heading1'],
        fontName='Helvetica-Bold',
        fontSize=20,
        leading=24,
        textColor=colors.HexColor('#1E3A8A'),
        alignment=1, # Center
        spaceAfter=4
    )
    subtitle_style = ParagraphStyle(
        'DocSub',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=10,
        leading=14,
        textColor=colors.HexColor('#64748B'),
        alignment=1,
        spaceAfter=15
    )
    sec_header_style = ParagraphStyle(
        'SecHeader',
        parent=styles['Heading2'],
        fontName='Helvetica-Bold',
        fontSize=12,
        leading=16,
        textColor=colors.HexColor('#FFFFFF')
    )
    body_style = ParagraphStyle(
        'Body',
        parent=styles['Normal'],
        fontName='Helvetica',
        fontSize=9,
        leading=13,
        textColor=colors.HexColor('#1E293B')
    )
    body_bold = ParagraphStyle(
        'BodyBold',
        parent=body_style,
        fontName='Helvetica-Bold'
    )

    story = []

    # Title & Header
    client_title = client['couple_name'] if client['couple_name'] else client['username']
    story.append(Paragraph(f"RINGKASAN DATA UNDANGAN: {client_title.upper()}", title_style))
    story.append(Paragraph(f"Status Pesanan: <b>{client['order_status']}</b> | Kelengkapan Data: <b>{wedding_data['completion_percent']}%</b> | Kontak: <b>{client['phone'] or '-'}</b>", subtitle_style))
    story.append(HRFlowable(width="100%", thickness=1.5, color=colors.HexColor('#33443C'), spaceAfter=12))

    def make_section_header(title):
        t = Table([[Paragraph(title, sec_header_style)]], colWidths=[520])
        t.setStyle(TableStyle([
            ('BACKGROUND', (0,0), (-1,-1), colors.HexColor('#33443C')),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
            ('LEFTPADDING', (0,0), (-1,-1), 8),
        ]))
        return t

    # 1. Tema Undangan
    story.append(make_section_header("1. TEMA & DESAIN UNDANGAN"))
    story.append(Spacer(1, 4))
    theme_name = wedding_data['theme_name'] if ('theme_name' in wedding_data.keys() and wedding_data['theme_name']) else (theme['name'] if theme else '-')
    theme_table_data = [
        [Paragraph("Nama Tema", body_bold), Paragraph(theme_name, body_style)]
    ]
    tt = Table(theme_table_data, colWidths=[120, 400])
    tt.setStyle(TableStyle([
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
        ('BACKGROUND', (0,0), (0,-1), colors.HexColor('#F8FAFC')),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(tt)
    story.append(Spacer(1, 10))

    # 2. Biodata Pengantin
    story.append(make_section_header("2. BIODATA MEMPELAI"))
    story.append(Spacer(1, 4))
    bio_table_data = [
        [Paragraph("<b>Detail</b>", body_style), Paragraph("<b>Mempelai Pria (Groom)</b>", body_style), Paragraph("<b>Mempelai Wanita (Bride)</b>", body_style)],
        [Paragraph("Nama Lengkap", body_bold), Paragraph(wedding_data['groom_name'] or '-', body_style), Paragraph(wedding_data['bride_name'] or '-', body_style)],
        [Paragraph("Nama Panggilan", body_bold), Paragraph(wedding_data['groom_nickname'] or '-', body_style), Paragraph(wedding_data['bride_nickname'] or '-', body_style)],
        [Paragraph("Urutan Anak", body_bold), Paragraph(wedding_data['groom_child_order'] or '-', body_style), Paragraph(wedding_data['bride_child_order'] or '-', body_style)],
        [Paragraph("Nama Ayah", body_bold), Paragraph(wedding_data['groom_father'] or '-', body_style), Paragraph(wedding_data['bride_father'] or '-', body_style)],
        [Paragraph("Nama Ibu", body_bold), Paragraph(wedding_data['groom_mother'] or '-', body_style), Paragraph(wedding_data['bride_mother'] or '-', body_style)],
        [Paragraph("Instagram", body_bold), Paragraph(wedding_data['groom_instagram'] or '-', body_style), Paragraph(wedding_data['bride_instagram'] or '-', body_style)]
    ]
    bt = Table(bio_table_data, colWidths=[100, 210, 210])
    bt.setStyle(TableStyle([
        ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
        ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#E2E8F0')),
        ('BACKGROUND', (0,1), (0,-1), colors.HexColor('#F8FAFC')),
        ('TOPPADDING', (0,0), (-1,-1), 4),
        ('BOTTOMPADDING', (0,0), (-1,-1), 4),
    ]))
    story.append(bt)
    story.append(Spacer(1, 10))

    # 3. Rangkaian Acara
    story.append(make_section_header("3. RANGKAIAN ACARA PERNIKAHAN"))
    story.append(Spacer(1, 4))
    events = json.loads(wedding_data['events_json'] or '[]')
    if events:
        ev_rows = [[Paragraph("<b>Acara</b>", body_style), Paragraph("<b>Waktu & Tanggal</b>", body_style), Paragraph("<b>Lokasi & Alamat</b>", body_style)]]
        for ev in events:
            ev_name = ev.get('name', 'Acara')
            ev_time = f"{ev.get('date', '')}<br/>{ev.get('start_time', '')} - {ev.get('end_time', '')}"
            ev_loc = f"<b>{ev.get('location_name', '')}</b><br/>{ev.get('address', '')}"
            if ev.get('maps_url'):
                ev_loc += f"<br/><font color='#2563EB'>Maps: {ev.get('maps_url')}</font>"
            ev_rows.append([Paragraph(ev_name, body_bold), Paragraph(ev_time, body_style), Paragraph(ev_loc, body_style)])
        et = Table(ev_rows, colWidths=[100, 140, 280])
        et.setStyle(TableStyle([
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
            ('BACKGROUND', (0,0), (-1,0), colors.HexColor('#E2E8F0')),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ]))
        story.append(et)
    else:
        story.append(Paragraph("<i>(Belum ada data acara diisi)</i>", body_style))
    story.append(Spacer(1, 10))

    # 4. Love Story
    story.append(make_section_header("4. KISAH CINTA (LOVE STORY)"))
    story.append(Spacer(1, 4))
    stories = json.loads(wedding_data['love_stories_json'] or '[]')
    if stories:
        st_rows = []
        for s in stories:
            st_rows.append([
                Paragraph(f"<b>{s.get('date_or_year', '-')}</b>", body_bold),
                Paragraph(f"<b>{s.get('title', '')}</b><br/>{s.get('story_text', '')}", body_style)
            ])
        st_t = Table(st_rows, colWidths=[100, 420])
        st_t.setStyle(TableStyle([
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
            ('BACKGROUND', (0,0), (0,-1), colors.HexColor('#F8FAFC')),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ]))
        story.append(st_t)
    else:
        story.append(Paragraph("<i>(Belum ada data kisah cinta)</i>", body_style))
    story.append(Spacer(1, 10))

    # 5. Kado Digital & Alamat
    story.append(make_section_header("5. KADO DIGITAL & PENGIRIMAN FISIK"))
    story.append(Spacer(1, 4))
    gifts = json.loads(wedding_data['gifts_json'] or '{}')
    gift_rows = []
    for b in gifts.get('bank_accounts', []):
        gift_rows.append([Paragraph("Rekening Bank", body_bold), Paragraph(f"<b>{b.get('bank_name')}</b> - No. {b.get('account_number')} (a.n. {b.get('account_holder')})", body_style)])
    for e in gifts.get('ewallets', []):
        gift_rows.append([Paragraph("E-Wallet", body_bold), Paragraph(f"<b>{e.get('provider')}</b> - No. {e.get('number')} (a.n. {e.get('account_holder')})", body_style)])
    phys = gifts.get('physical_gift', {})
    if phys.get('address'):
        gift_rows.append([Paragraph("Alamat Kado Fisik", body_bold), Paragraph(f"Penerima: <b>{phys.get('recipient')}</b> ({phys.get('phone')})<br/>{phys.get('address')}", body_style)])
    
    if gift_rows:
        gt = Table(gift_rows, colWidths=[120, 400])
        gt.setStyle(TableStyle([
            ('GRID', (0,0), (-1,-1), 0.5, colors.HexColor('#CBD5E1')),
            ('BACKGROUND', (0,0), (0,-1), colors.HexColor('#F8FAFC')),
            ('TOPPADDING', (0,0), (-1,-1), 4),
            ('BOTTOMPADDING', (0,0), (-1,-1), 4),
        ]))
        story.append(gt)
    else:
        story.append(Paragraph("<i>(Belum ada data kado digital)</i>", body_style))
    story.append(Spacer(1, 10))

    # 6. Turut Mengundang
    story.append(make_section_header("6. TURUT MENGUNDANG"))
    story.append(Spacer(1, 4))
    inviters = json.loads(wedding_data['guest_inviters_json'] or '[]')
    if inviters:
        inv_text = "<br/>".join([f"• {inv.get('name')}" for inv in inviters if inv.get('name')])
        story.append(Paragraph(inv_text, body_style))
    else:
        story.append(Paragraph("<i>(Belum ada nama yang turut mengundang)</i>", body_style))
    story.append(Spacer(1, 10))

    # 7. Backsound
    story.append(make_section_header("7. BACKSOUND MUSIK"))
    story.append(Spacer(1, 4))
    music_info = f"<b>Tipe:</b> {'Preset Pilihan Admin' if wedding_data['music_choice_type'] == 'preset' else 'File Musik Unggahan Sendiri'}<br/>"
    music_info += f"<b>Lagu:</b> {wedding_data['music_title'] or '-'} | <b>Artis:</b> {wedding_data['music_artist'] or '-'}"
    story.append(Paragraph(music_info, body_style))

    # Admin internal note (if any, footer)
    if client['notes_internal']:
        story.append(Spacer(1, 10))
        story.append(Paragraph(f"<b>Catatan Internal Admin:</b> <i>{client['notes_internal']}</i>", ParagraphStyle('Notes', parent=body_style, textColor=colors.HexColor('#9A3412'))))

    doc.build(story)
    buf.seek(0)
    return buf.getvalue()


def generate_zip_bundle(client, wedding_data, theme, base_dir):
    """Packages all data, Excel, PDF, JSON, and uploaded media files into a complete ZIP archive."""
    excel_bytes = generate_excel_bytes(client, wedding_data, theme)
    pdf_bytes = generate_pdf_bytes(client, wedding_data, theme, base_dir)

    safe_name = "".join([c if c.isalnum() else "_" for c in client['couple_name']]).strip("_")
    
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, 'w', zipfile.ZIP_DEFLATED) as z:
        # Add Excel
        z.writestr(f"{safe_name}_data_undangan.xlsx", excel_bytes)
        # Add PDF
        z.writestr(f"{safe_name}_ringkasan_biodata.pdf", pdf_bytes)
        # Add JSON
        full_json = {
            'client': dict(client),
            'wedding_data': dict(wedding_data),
            'theme': dict(theme) if theme else None
        }
        z.writestr(f"{safe_name}_data_undangan.json", json.dumps(full_json, indent=2, default=str))

        # Add Media files if they exist on disk
        media_paths = []
        if wedding_data['groom_photo']:
            media_paths.append(('media/foto_mempelai_pria' + os.path.splitext(wedding_data['groom_photo'])[1], wedding_data['groom_photo']))
        if wedding_data['bride_photo']:
            media_paths.append(('media/foto_mempelai_wanita' + os.path.splitext(wedding_data['bride_photo'])[1], wedding_data['bride_photo']))

        # Love stories photos
        stories = json.loads(wedding_data['love_stories_json'] or '[]')
        for idx, s in enumerate(stories, 1):
            if s.get('photo'):
                media_paths.append((f"media/love_story_{idx}" + os.path.splitext(s['photo'])[1], s['photo']))

        # Custom music
        if wedding_data['music_custom_file']:
            media_paths.append(('media/backsound_kustom' + os.path.splitext(wedding_data['music_custom_file'])[1], wedding_data['music_custom_file']))

        for zip_target, rel_path in media_paths:
            # rel_path can be /static/uploads/photos/... or uploads/photos/...
            clean_rel = rel_path.lstrip('/')
            abs_path = os.path.join(base_dir, clean_rel)
            if os.path.exists(abs_path):
                z.write(abs_path, zip_target)

    buf.seek(0)
    return buf.getvalue()
